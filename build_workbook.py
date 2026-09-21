"""Builds the Excel workbook behind the dashboard.

SYNTHETIC DATA. Demonstrates a method, not any company's business.

The point of this file is lineage. It lays out the raw order book as an ERP would
hand it over, then every derived field and aggregate as a live Excel formula against
that raw tab, so each number on the dashboard can be traced back to a source row.

Anything Excel can reasonably compute is a formula. The Monte Carlo draw is the one step
it cannot, so the tabs that carry simulated percentiles are marked as model output.

Output: SOP_data_and_calculations.xlsx
"""

from __future__ import annotations

import os
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula


def pctile(ws, row, col, cond, values, q):
    """Conditional percentile as a genuine CSE array formula.

    Two traps here. openpyxl writes post-2007 function names without the _xlfn.
    prefix Excel needs, so PERCENTILE.INC comes back #NAME?; the legacy
    PERCENTILE has no such problem. And PERCENTILE(IF(...)) only evaluates as
    an array formula, which has to be declared rather than written as text.
    """
    ref = f"{get_column_letter(col)}{row}"
    ws[ref] = ArrayFormula(ref, f"=PERCENTILE(IF({cond},{values}),{q})")
    ws[ref].number_format = "0.00"

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "data")
OUT = os.path.join(HERE, "SOP_data_and_calculations.xlsx")

ORANGE = "FE5000"
DARK = "111827"
GREY = "F5F5F5"
BORDER = "D4D4D4"

MATERIAL_RATE = 0.50
REGIONS = ["North America", "EMEA", "APAC", "LATAM"]
MARKETS = ["US / Canada", "Europe", "MENA", "Australia/SEA", "India", "LATAM"]

hdr_font = Font(bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor=DARK)
lab = Font(bold=True, size=10)
thin = Side(style="thin", color=BORDER)
box = Border(bottom=thin)
note_font = Font(italic=True, size=9, color="737373")
title_font = Font(bold=True, size=13, color=DARK)
sec_font = Font(bold=True, size=10, color=ORANGE)


def sheet(wb, name):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    return ws


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols):
        cell = ws.cell(row=row, column=1 + i, value=c)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 30
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(1 + i)].width = w


def title(ws, text):
    ws["A1"] = text
    ws["A1"].font = title_font


def main() -> None:
    orders = pd.read_csv(os.path.join(DATA, "orders.csv"),
                         parse_dates=["booked_month", "expected_conversion",
                                      "actual_conversion"])
    mc = pd.read_csv(os.path.join(DATA, "mc_monthly.csv"), parse_dates=["month"])
    cov = pd.read_csv(os.path.join(DATA, "mc_coverage.csv"),
                      parse_dates=["cutoff", "month"])

    orders = orders.sort_values("order_id").reset_index(drop=True)
    n = len(orders)
    months = sorted(orders.booked_month.unique())

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---------------------------------------------------------------- Contents
    ws = sheet(wb, "Contents")
    title(ws, "S&OP demand and material review - data and calculations")
    ws["A3"] = ("SYNTHETIC DATA. Region names and the North America revenue weighting come from "
                "public filings and public role titles. Every order, date and dollar figure is "
                "generated. This demonstrates a method and represents no company's actual business.")
    ws["A3"].font = Font(italic=True, size=10, color="B53A00")
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A3:F5")
    rows = [
        ("Raw Orders", "The order book as an ERP would extract it. Values only, one row per project.", "source"),
        ("Engineered Orders", "Fields derived per order: material committed, slip, status, material at risk.", "formulas"),
        ("Monthly by Region", "Bookings, conversions, backlog and book-to-bill by region and month.", "formulas"),
        ("Slip by Market", "Schedule slip distribution per market, P10 / P50 / P90.", "formulas"),
        ("Region vs Market", "Each market's P90 against its region's blended P90.", "formulas"),
        ("Lookback", "Expected against actual conversion by region and month, with cumulative bias.", "formulas"),
        ("Commitment", "Material committed against backlog at a range of realization assumptions.", "formulas"),
        ("MC Forward", "Forward requirement by month. Simulated P10/P50/P90; spread and firm share are formulas.", "mixed"),
        ("MC Coverage", "Walk-forward test. How often the stated interval contained the outcome, by months ahead.", "mixed"),
    ]
    header(ws, 7, ["Tab", "What is on it", "Type"], [22, 86, 15])
    for i, (a, b, c) in enumerate(rows):
        r = 8 + i
        ws.cell(row=r, column=1, value=a).font = lab
        ws.cell(row=r, column=2, value=b)
        cc = ws.cell(row=r, column=3, value=c)
        cc.alignment = Alignment(horizontal="center")
        if c != "formulas":
            cc.font = Font(size=10, color="B53A00", bold=True)
        for j in range(1, 4):
            ws.cell(row=r, column=j).border = box
    r = 8 + len(rows) + 2
    ws.cell(row=r, column=1, value="On the two model-output tabs").font = sec_font
    ws.cell(row=r + 1, column=1, value=(
        "Excel reproduces every aggregate in this workbook, including the whole coverage test. The "
        "one thing it cannot reasonably do is run 4,000 simulations across the order book, so the "
        "P10/P50/P90 columns on the two MC tabs are simulated values. Everything derived from them "
        "is a formula. The scripts are in the same repository."))
    ws.cell(row=r + 1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 3, end_column=6)
    ws.cell(row=r + 5, column=1, value="Material rate used throughout (from the case study)").font = lab
    ws.cell(row=r + 5, column=3, value=MATERIAL_RATE).number_format = "0%"

    # ---------------------------------------------------------------- Raw
    ws = sheet(wb, "Raw Orders")
    title(ws, "Raw order book")
    cols = ["order_id", "region", "market", "booked_month", "mw", "order_value",
            "promised_months", "expected_conversion", "actual_conversion",
            "cancelled", "slip_months"]
    header(ws, 4, cols, [12, 15, 15, 14, 9, 14, 11, 18, 18, 11, 12])
    for i, row in orders.iterrows():
        r = 5 + i
        vals = [row.order_id, row.region, row.market, row.booked_month.date(),
                row.mw, row.order_value, int(row.promised_months),
                row.expected_conversion.date(),
                None if pd.isna(row.actual_conversion) else row.actual_conversion.date(),
                bool(row.cancelled), float(row.slip_months)]
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=1 + j, value=v)
            if j in (3, 7, 8):
                c.number_format = "yyyy-mm-dd"
            elif j == 5:
                c.number_format = "#,##0"
            elif j in (4, 10):
                c.number_format = "#,##0.0"
    ws.freeze_panes = "A5"
    last = 4 + n

    # ---------------------------------------------------------------- Engineered
    ws = sheet(wb, "Engineered Orders")
    title(ws, "Engineered fields")
    cols = ["order_id", "region", "market", "mw", "order_value",
            "material_committed", "expected_conversion", "actual_conversion",
            "slip_months", "slip_check_from_dates", "status", "material_at_risk"]
    header(ws, 4, cols, [12, 15, 15, 9, 14, 16, 18, 18, 12, 18, 13, 15])
    ws.cell(row=3, column=6, value="= order_value x material rate").font = note_font
    ws.cell(row=3, column=10, value="= whole months between expected and actual, "
                                    "a check on the supplied slip").font = note_font
    for i in range(n):
        r = 5 + i
        src = 5 + i
        f = {
            1: f"='Raw Orders'!A{src}",
            2: f"='Raw Orders'!B{src}",
            3: f"='Raw Orders'!C{src}",
            4: f"='Raw Orders'!E{src}",
            5: f"='Raw Orders'!F{src}",
            6: f"='Raw Orders'!F{src}*{MATERIAL_RATE}",
            7: f"='Raw Orders'!H{src}",
            8: f"='Raw Orders'!I{src}",
            9: f"='Raw Orders'!K{src}",
            10: (f"=IF('Raw Orders'!J{src},\"\","
                 f"(YEAR('Raw Orders'!I{src})-YEAR('Raw Orders'!H{src}))*12"
                 f"+MONTH('Raw Orders'!I{src})-MONTH('Raw Orders'!H{src}))"),
            11: (f"=IF('Raw Orders'!J{src},\"Cancelled\","
                 f"IF('Raw Orders'!K{src}>=2,\"Slipped\",\"On plan\"))"),
            12: f"=IF(K{r}=\"On plan\",0,F{r})",
        }
        for col, formula in f.items():
            c = ws.cell(row=r, column=col, value=formula)
            if col in (5, 6, 12):
                c.number_format = "#,##0"
            elif col in (4, 9, 10):
                c.number_format = "#,##0.0"
            elif col in (7, 8):
                c.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A5"

    RAW = "'Raw Orders'"
    ENG = "'Engineered Orders'"
    R_REG = f"{RAW}!$B$5:$B${last}"
    R_MKT = f"{RAW}!$C$5:$C${last}"
    R_BOOK = f"{RAW}!$D$5:$D${last}"
    R_MW = f"{RAW}!$E$5:$E${last}"
    R_VAL = f"{RAW}!$F$5:$F${last}"
    R_EXP = f"{RAW}!$H$5:$H${last}"
    R_ACT = f"{RAW}!$I$5:$I${last}"
    R_CAN = f"{RAW}!$J$5:$J${last}"
    R_SLIP = f"{RAW}!$K$5:$K${last}"
    E_MAT = f"{ENG}!$F$5:$F${last}"

    # ---------------------------------------------------------------- Monthly
    ws = sheet(wb, "Monthly by Region")
    title(ws, "Monthly aggregates by region")
    cols = ["region", "month", "orders booked", "MW booked", "value booked ($)",
            "value converted ($)", "backlog ($)", "book-to-bill",
            "material vs backlog ($)"]
    header(ws, 4, cols, [15, 12, 13, 12, 16, 17, 15, 12, 19])
    r = 5
    for region in REGIONS:
        for m in months:
            d = pd.Timestamp(m).date()
            ws.cell(row=r, column=1, value=region)
            c = ws.cell(row=r, column=2, value=d)
            c.number_format = "yyyy-mm"
            ws.cell(row=r, column=3,
                    value=f"=COUNTIFS({R_REG},$A{r},{R_BOOK},$B{r})").number_format = "#,##0"
            ws.cell(row=r, column=4,
                    value=f"=SUMIFS({R_MW},{R_REG},$A{r},{R_BOOK},$B{r})").number_format = "#,##0"
            ws.cell(row=r, column=5,
                    value=f"=SUMIFS({R_VAL},{R_REG},$A{r},{R_BOOK},$B{r})").number_format = "#,##0"
            ws.cell(row=r, column=6,
                    value=(f"=SUMIFS({R_VAL},{R_REG},$A{r},{R_ACT},$B{r},"
                           f"{R_CAN},FALSE)")).number_format = "#,##0"
            ws.cell(row=r, column=7,
                    value=(f"=SUMIFS({R_VAL},{R_REG},$A{r},{R_BOOK},\"<=\"&$B{r},{R_CAN},FALSE)"
                           f"-SUMIFS({R_VAL},{R_REG},$A{r},{R_ACT},\"<=\"&$B{r},"
                           f"{R_CAN},FALSE)")).number_format = "#,##0"
            ws.cell(row=r, column=8,
                    value=f"=IF(F{r}=0,\"\",E{r}/F{r})").number_format = "0.00"
            ws.cell(row=r, column=9,
                    value=f"=G{r}*{MATERIAL_RATE}").number_format = "#,##0"
            for j in range(1, 10):
                ws.cell(row=r, column=j).border = box
            r += 1
    ws.freeze_panes = "A5"

    # ---------------------------------------------------------------- Slip
    ws = sheet(wb, "Slip by Market")
    title(ws, "Schedule slip distribution by market")
    header(ws, 4, ["market", "region", "orders", "P10", "median", "P90", "mean"],
           [16, 16, 10, 10, 10, 10, 10])
    mkt_region = orders.drop_duplicates("market").set_index("market").region.to_dict()
    for i, mk in enumerate(MARKETS):
        r = 5 + i
        ws.cell(row=r, column=1, value=mk).font = lab
        ws.cell(row=r, column=2, value=mkt_region.get(mk, ""))
        ws.cell(row=r, column=3,
                value=f"=COUNTIFS({R_MKT},$A{r},{R_CAN},FALSE)").number_format = "#,##0"
        for j, q in ((4, 0.10), (5, 0.50), (6, 0.90)):
            pctile(ws, r, j, f"({R_MKT}=$A{r})*({R_CAN}=FALSE)", R_SLIP, q)
        ws.cell(row=r, column=7, value=(
            f"=AVERAGEIFS({R_SLIP},{R_MKT},$A{r},{R_CAN},FALSE)")).number_format = "0.00"
        for j in range(1, 8):
            ws.cell(row=r, column=j).border = box
    rr = 5 + len(MARKETS) + 2
    ws.cell(row=rr, column=1, value=(
        "The percentile columns are array formulas over the raw slip column, filtered to the "
        "market. Written as CSE arrays so they evaluate in any Excel version.")).font = note_font

    # ---------------------------------------------------------------- Blend
    ws = sheet(wb, "Region vs Market")
    title(ws, "Market against regional blend")
    header(ws, 4, ["region", "region P90", "market", "market P90", "gap vs region"],
           [16, 12, 16, 12, 14])
    r = 5
    for region in REGIONS:
        mks = [m for m in MARKETS if mkt_region.get(m) == region]
        for mk in mks:
            ws.cell(row=r, column=1, value=region)
            pctile(ws, r, 2, f"({R_REG}=$A{r})*({R_CAN}=FALSE)", R_SLIP, 0.9)
            ws.cell(row=r, column=3, value=mk)
            pctile(ws, r, 4, f"({R_MKT}=$C{r})*({R_CAN}=FALSE)", R_SLIP, 0.9)
            ws.cell(row=r, column=5, value=f"=D{r}-B{r}").number_format = "+0.00;-0.00"
            for j in range(1, 6):
                ws.cell(row=r, column=j).border = box
            r += 1

    # ---------------------------------------------------------------- Lookback
    ws = sheet(wb, "Lookback")
    title(ws, "Conversion lookback")
    header(ws, 4, ["region", "month", "expected MW", "actual MW", "realization",
                   "bias MW", "cumulative bias MW"],
           [15, 12, 13, 12, 12, 12, 18])
    r = 5
    for region in REGIONS:
        first = r
        for m in months:
            d = pd.Timestamp(m).date()
            ws.cell(row=r, column=1, value=region)
            ws.cell(row=r, column=2, value=d).number_format = "yyyy-mm"
            ws.cell(row=r, column=3,
                    value=f"=SUMIFS({R_MW},{R_REG},$A{r},{R_EXP},$B{r})").number_format = "#,##0"
            ws.cell(row=r, column=4,
                    value=(f"=SUMIFS({R_MW},{R_REG},$A{r},{R_ACT},$B{r},"
                           f"{R_CAN},FALSE)")).number_format = "#,##0"
            ws.cell(row=r, column=5,
                    value=f"=IF(C{r}=0,\"\",D{r}/C{r})").number_format = "0%"
            ws.cell(row=r, column=6, value=f"=D{r}-C{r}").number_format = "#,##0"
            ws.cell(row=r, column=7,
                    value=f"=SUM($F${first}:$F{r})").number_format = "#,##0"
            for j in range(1, 8):
                ws.cell(row=r, column=j).border = box
            r += 1
    ws.freeze_panes = "A5"

    # ---------------------------------------------------------------- Commitment
    ws = sheet(wb, "Commitment")
    title(ws, "Material committed against backlog")
    rates = [1.00, 0.95, 0.90, 0.85, 0.80, 0.75, 0.70]
    header(ws, 4, ["region", "closing backlog ($)"] + [f"at {int(x*100)}%" for x in rates]
           + ["exposure 100% vs 70%"],
           [16, 19] + [12] * len(rates) + [20])
    lastm = pd.Timestamp(months[-1]).date()
    r = 5
    for region in REGIONS:
        ws.cell(row=r, column=1, value=region).font = lab
        ws.cell(row=r, column=2, value=(
            f"=SUMIFS({R_VAL},{R_REG},$A{r},{R_BOOK},\"<=\"&DATE({lastm.year},{lastm.month},1),"
            f"{R_CAN},FALSE)"
            f"-SUMIFS({R_VAL},{R_REG},$A{r},{R_ACT},\"<=\"&DATE({lastm.year},{lastm.month},1),"
            f"{R_CAN},FALSE)")).number_format = "#,##0"
        for j, rate in enumerate(rates):
            ws.cell(row=r, column=3 + j,
                    value=f"=$B{r}*{MATERIAL_RATE}*{rate}").number_format = "#,##0"
        ws.cell(row=r, column=3 + len(rates),
                value=f"=C{r}-{get_column_letter(2+len(rates))}{r}").number_format = "#,##0"
        for j in range(1, 4 + len(rates)):
            ws.cell(row=r, column=j).border = box
        r += 1
    tot = r
    ws.cell(row=tot, column=1, value="Total").font = lab
    for j in range(2, 4 + len(rates)):
        L = get_column_letter(j)
        ws.cell(row=tot, column=j,
                value=f"=SUM({L}5:{L}{tot-1})").number_format = "#,##0"
        ws.cell(row=tot, column=j).font = lab

    # ---------------------------------------------------------------- MC out
    ws = sheet(wb, "MC Forward")
    title(ws, "Forward material requirement")
    t = mc[(mc.scope == "TOTAL")].pivot_table(index="month", columns="layer",
                                              values="p50").reset_index()
    band = mc[(mc.scope == "TOTAL") & (mc.layer == "all")].sort_values("month")
    header(ws, 4, ["month", "already booked ($)", "not yet booked ($)",
                   "expected total ($)", "low (P10)", "high (P90)",
                   "spread", "firm share"],
           [12, 18, 18, 17, 14, 14, 12, 12])
    ws.cell(row=3, column=7, value="= high - low, the width of the stated 80% interval"
            ).font = note_font
    for i in range(len(band)):
        r = 5 + i
        b = band.iloc[i]
        row_t = t[t.month == b.month].iloc[0]
        ws.cell(row=r, column=1, value=b.month.date()).number_format = "yyyy-mm"
        for j, v in enumerate([row_t.get("firm", 0), row_t.get("pipeline", 0),
                               b.p50, b.p10, b.p90]):
            c = ws.cell(row=r, column=2 + j,
                        value=None if v is None or pd.isna(v) else float(v))
            c.number_format = "#,##0"
        ws.cell(row=r, column=7, value=f"=F{r}-E{r}").number_format = "#,##0"
        ws.cell(row=r, column=8, value=f"=IF(D{r}=0,\"\",B{r}/D{r})").number_format = "0%"
        for j in range(1, 9):
            ws.cell(row=r, column=j).border = box
    nb = 4 + len(band)
    ws.cell(row=nb + 2, column=1, value=(
        "Already booked and not yet booked are medians of two separate simulations, so they do "
        "not sum exactly to the expected total, which is the median of the combined draw. "
        "Medians are not additive.")).font = note_font
    ws.merge_cells(start_row=nb + 2, start_column=1, end_row=nb + 3, end_column=8)

    ws = sheet(wb, "MC Coverage")
    title(ws, "Walk-forward interval coverage")
    nc = len(cov)
    lastr = 4 + nc
    ws.cell(row=2, column=1, value=(
        "Each row re-runs the model at a past cutoff, using only what was known at that cutoff, "
        "and scores the forecast against what the month turned out to be.")).font = note_font
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=9)
    header(ws, 4, ["cutoff", "month", "months ahead", "actual ($)", "P10", "P25",
                   "P50", "P75", "P90", "spread (P90-P10)",
                   "inside 80%", "inside 50%"],
           [12, 12, 13, 14, 13, 13, 13, 13, 13, 16, 12, 12])

    for i in range(nc):
        r = 5 + i
        c0 = cov.iloc[i]
        ws.cell(row=r, column=1, value=c0.cutoff.date()).number_format = "yyyy-mm"
        ws.cell(row=r, column=2, value=c0.month.date()).number_format = "yyyy-mm"
        ws.cell(row=r, column=3, value=int(c0.offset))
        for j, v in enumerate([c0.actual, c0.p10, c0.p25, c0.p50, c0.p75, c0.p90]):
            ws.cell(row=r, column=4 + j, value=float(v)).number_format = "#,##0"
        ws.cell(row=r, column=10, value=f"=I{r}-E{r}").number_format = "#,##0"
        ws.cell(row=r, column=11, value=f"=IF(AND(D{r}>=E{r},D{r}<=I{r}),1,0)")
        ws.cell(row=r, column=12, value=f"=IF(AND(D{r}>=F{r},D{r}<=H{r}),1,0)")
        for j in range(1, 13):
            ws.cell(row=r, column=j).border = box

    IN80 = f"$K$5:$K${lastr}"
    IN50 = f"$L$5:$L${lastr}"
    AH = f"$C$5:$C${lastr}"
    MO = f"$B$5:$B${lastr}"

    s = lastr + 3
    ws.cell(row=s, column=1, value="What the test measures").font = sec_font
    ws.cell(row=s + 1, column=1, value=(
        "%d cutoffs, six months forward from each. That is %d scored forecasts but only %d "
        "distinct outcome months, because a six-month window scores the same month from as many "
        "as six cutoffs. Everything below is measurement. No multiplier is fitted here and none "
        "is applied to the forward tab: conformal calibration assumes the scores are "
        "exchangeable, scores sharing an outcome month are not, and without that the guarantee "
        "fails at any sample size. This also scores the booked backlog only, since pipeline "
        "demand at a past cutoff cannot be checked against the order book."
        % (cov.cutoff.nunique(), len(cov), cov.month.nunique()))).font = note_font
    ws.merge_cells(start_row=s + 1, start_column=1, end_row=s + 2, end_column=9)

    rows_sum = [
        ("forecasts scored (n)", f"=COUNT({IN80})", "#,##0"),
        ("distinct outcome months",
         f'=SUMPRODUCT(({MO}<>"")/COUNTIF({MO},{MO}&""))', "#,##0"),
        ("stated 80% interval, coverage achieved", f"=AVERAGE({IN80})", "0%"),
        ("stated 50% interval, coverage achieved", f"=AVERAGE({IN50})", "0%"),
        ("misses above the interval",
         f"=SUMPRODUCT(--($D$5:$D${lastr}>$I$5:$I${lastr}))", "#,##0"),
        ("misses below the interval",
         f"=SUMPRODUCT(--($D$5:$D${lastr}<$E$5:$E${lastr}))", "#,##0"),
    ]
    for i, (labeltext, formula, fmt) in enumerate(rows_sum):
        r = s + 4 + i
        ws.cell(row=r, column=1, value=labeltext).font = lab
        ws[f"C{r}"] = formula
        ws[f"C{r}"].number_format = fmt
        if "coverage achieved" in labeltext:
            ws.cell(row=r, column=1).font = Font(bold=True, size=10, color=ORANGE)

    h = s + 4 + len(rows_sum) + 2
    ws.cell(row=h, column=1, value="Coverage by how far ahead the call was made").font = sec_font
    ws.cell(row=h + 1, column=1, value=(
        "The interval is weakest one month out and holds further out, which is the reverse of "
        "the usual shape. The cause is not scatter. Sorted by cutoff, the month +1 error runs in "
        "one direction and then the other, shown in the second table: the early cutoffs forecast "
        "high and the recent ones forecast low. That is drift in the front month, which a "
        "tracking signal catches and a wider interval only masks.")).font = note_font
    ws.merge_cells(start_row=h + 1, start_column=1, end_row=h + 2, end_column=9)
    header(ws, h + 4, ["months ahead", "forecasts", "inside 80%"])
    for i, off in enumerate(sorted(cov.offset.unique())):
        r = h + 5 + i
        ws.cell(row=r, column=1, value=int(off))
        ws.cell(row=r, column=2, value=f"=COUNTIFS({AH},A{r})").number_format = "#,##0"
        # 0.0% not 0%: 62.5 formatted as 0% shows 63 in Excel and 62 on the page,
        # because the two round half-values in opposite directions
        ws.cell(row=r, column=3,
                value=f"=AVERAGEIFS({IN80},{AH},A{r})").number_format = "0.0%"
        for j in range(1, 4):
            ws.cell(row=r, column=j).border = box

    # the drift itself, so an Excel reader can see it rather than take it on trust
    worst = int(cov.groupby("offset").in_80.mean().idxmin())
    w = cov[cov.offset == worst].sort_values("cutoff")
    d = h + 5 + len(cov.offset.unique()) + 2
    ws.cell(row=d, column=1,
            value=f"Month +{worst} error by cutoff, in order").font = sec_font
    header(ws, d + 2, ["cutoff", "actual ($)", "P50 ($)", "error ($)", "direction"])
    for i in range(len(w)):
        r = d + 3 + i
        row = w.iloc[i]
        src = 5 + int(cov.index.get_loc(row.name))
        ws.cell(row=r, column=1, value=row.cutoff.date()).number_format = "yyyy-mm"
        ws.cell(row=r, column=2, value=f"=D{src}").number_format = "#,##0"
        ws.cell(row=r, column=3, value=f"=G{src}").number_format = "#,##0"
        ws.cell(row=r, column=4, value=f"=B{r}-C{r}").number_format = "+#,##0;-#,##0"
        ws.cell(row=r, column=5,
                value=f'=IF(D{r}<0,"forecast high","forecast low")')
        for j in range(1, 6):
            ws.cell(row=r, column=j).border = box

    wb.save(OUT)
    print("wrote", OUT)
    print("orders: %d   sheets: %d" % (n, len(wb.sheetnames)))
    print("sheets:", " | ".join(wb.sheetnames))


if __name__ == "__main__":
    main()
